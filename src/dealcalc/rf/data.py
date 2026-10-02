"""Offline import and normalization for Russian marketplace listings.

The importer reads user-provided CSV/JSON/JSONL and Excel (.xlsx) files
only. It deliberately does not make network requests. This makes it suitable
for cached exports and keeps collection policy separate from valuation
formulas.

Listing types: ``property`` and ``vehicle`` (sale offers), ``rent`` (rent of
premises, annualized), ``income`` (price with NOI or gross income, for the
capitalization rate and GRM), ``business`` (value and metric of a company,
for multiples) and ``machinery``.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Tuple


LISTING_TYPES = {"property", "vehicle", "rent", "income", "business", "machinery"}

_ALIASES = {
    "source": ("source", "источник", "площадка", "сайт"),
    "listing_id": ("listing_id", "id", "идентификатор", "номер объявления"),
    "url": ("url", "link", "ссылка", "ссылка на объявление"),
    "address": ("address", "адрес", "местоположение"),
    "cadastral_number": ("cadastral_number", "кадастровый номер", "кадастровый"),
    "vin": ("vin", "вин", "vin-номер"),
    "date": (
        "date",
        "дата публикации",
        "дата объявления",
        "дата размещения",
        "размещено",
    ),
    "price_type": ("price_type", "тип цены"),
    "collected_at": ("collected_at", "дата сбора", "дата_сбора"),
    "region": ("region", "регион", "область", "край"),
    "city": ("city", "город", "населенный пункт", "населённый пункт"),
    "price_rub": ("price_rub", "price", "цена", "стоимость", "цена руб"),
    "area_sqm": ("area_sqm", "area", "площадь", "площадь м2", "площадь м²"),
    "rooms": ("rooms", "комнаты", "комнат"),
    "floor": ("floor", "этаж"),
    "total_floors": ("total_floors", "этажей", "всего этажей"),
    "year": ("year", "год", "год выпуска", "год постройки"),
    "condition": ("condition", "состояние"),
    "conditions": ("conditions", "условия", "условия сделки", "условия продажи"),
    "reliability": ("reliability", "надёжность", "надежность", "достоверность"),
    "brand": ("brand", "марка", "make"),
    "model": ("model", "модель"),
    "mileage_km": ("mileage_km", "mileage", "пробег", "пробег км"),
    "engine_power_hp": ("engine_power_hp", "мощность", "мощность лс", "мощность л.с."),
    "transmission": ("transmission", "коробка", "коробка передач"),
    "drive": ("drive", "привод"),
    "adjustment_pct": ("adjustment_pct", "корректировка", "корректировка %"),
    "weight": ("weight", "вес"),
    "adjustments": ("adjustments", "корректировки"),
    # Rent.
    "rent_rub": ("rent_rub", "rent", "арендная плата", "ставка аренды", "аренда"),
    "rent_sqm": (
        "rent_sqm", "ставка аренды за м2", "ставка аренды за м²", "ставка за м2", "ставка за м²",
        "аренда за м2", "аренда за м²",
    ),
    "rent_period": ("rent_period", "период", "период оплаты", "период аренды"),
    # Income properties.
    "noi": ("noi", "чод", "чистый операционный доход"),
    "gross_income": ("gross_income", "валовой доход", "годовой валовой доход"),
    # Business.
    "value": ("value", "стоимость компании", "стоимость бизнеса", "капитализация", "ev", "цена сделки"),
    "metric": ("metric", "показатель", "значение показателя"),
    "name": ("name", "компания", "наименование", "название"),
    "industry": ("industry", "отрасль"),
    # Machinery.
    "operating_hours": ("operating_hours", "наработка", "моточасы", "наработка моточасов"),
}

_RENT_PERIODS = {
    "month": "month", "monthly": "month", "месяц": "month", "мес": "month", "мес.": "month",
    "в месяц": "month", "ежемесячно": "month",
    "year": "year", "annual": "year", "год": "year", "в год": "year", "ежегодно": "year",
}

# Fields read only for a given listing type.
_TYPE_TEXT_FIELDS = {"business": ("name", "industry"), "machinery": ("name",)}
_TYPE_NUMBER_FIELDS = {
    "income": ("noi", "gross_income"),
    "business": ("value", "metric"),
    "machinery": ("operating_hours",),
}

_TEXT_FIELDS = {
    "listing_id",
    "url",
    "address",
    "cadastral_number",
    "vin",
    "date",
    "collected_at",
    "region",
    "city",
    "condition",
    "conditions",
    "reliability",
    "brand",
    "model",
    "transmission",
    "drive",
}
_NUMBER_FIELDS = {
    "price_rub",
    "area_sqm",
    "rooms",
    "floor",
    "total_floors",
    "year",
    "mileage_km",
    "engine_power_hp",
}


def _lookup(row: Mapping[str, Any], field: str) -> Any:
    """Find a field using canonical and common Russian aliases."""

    lowered = {str(key).strip().lower(): value for key, value in row.items()}
    for alias in _ALIASES[field]:
        if alias.lower() in lowered:
            return lowered[alias.lower()]
    return None


# A number, an optional multiplier and an optional unit: "5 млн руб.", "45 м2".
_NUMBER_TEXT = re.compile(
    r"^(?P<number>[-+]?[0-9][0-9 .,]*?)\s*"
    r"(?:(?P<multiplier>тыс|млн|млрд)\.?)?\s*"
    r"(?:руб(?:лей|ля|ль)?\.?|р\.|₽|rub|км|km|м2|м²|кв\.?\s*м\.?|sqm|л\.?\s*с\.?|hp|%)?$"
)
_MULTIPLIERS = {"тыс": Decimal(1_000), "млн": Decimal(1_000_000), "млрд": Decimal(1_000_000_000)}


def parse_number(value: Any, field: str) -> Optional[float]:
    """Parse a number from a CSV/JSON value, including Russian price strings.

    Thousand separators, a decimal comma, the multipliers "тыс.", "млн",
    "млрд" and a trailing unit (руб., ₽, км, м2, л.с., %) are understood.
    Any other letters ("1e6", "USD", "от 5 000 000") are rejected rather than
    dropped, so the number never changes its scale silently.
    """

    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"{field} must be a number")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace("\u00a0", " ")
        if not text:
            return None
        match = _NUMBER_TEXT.match(text.lower())
        if not match:
            raise ValueError(f"{field} must be a number, got {text!r}")
        multiplier = _MULTIPLIERS.get(match.group("multiplier") or "", Decimal(1))
        text = match.group("number").replace(" ", "").rstrip(".,")
        if "," in text and "." in text:
            # The last separator is the decimal one: 1.200.000,50 or 1,200,000.50.
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif text.count(",") == 1:
            text = text.replace(",", ".")
        elif text.count(",") > 1:
            text = text.replace(",", "")
        elif text.count(".") > 1:
            text = text.replace(".", "")
        try:
            number = float(Decimal(text) * multiplier)
        except InvalidOperation as exc:
            raise ValueError(f"{field} must be a number") from exc
    if not math.isfinite(number):
        raise ValueError(f"{field} must be finite")
    return number


def _parse_adjustments(value: Any) -> Optional[List[Dict[str, Any]]]:
    """Read step adjustments: a list, or a JSON list in a CSV cell."""

    if value is None or value == "":
        return None
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError("adjustments must be a JSON list of objects") from exc
    if not isinstance(value, list) or not all(isinstance(step, Mapping) for step in value):
        raise ValueError("adjustments must be a JSON list of objects")
    return [dict(step) for step in value]


_PRICE_TYPES = {
    "сделка": "сделка",
    "transaction": "сделка",
    "цена сделки": "сделка",
    "предложение": "предложение",
    "offer": "предложение",
    "цена предложения": "предложение",
}


def _price_type(value: Any) -> Tuple[str, Optional[str]]:
    """Marketplace exports are offers unless the file says otherwise.

    An unknown value ("Продажа", "Аренда" — often the deal type) is left
    empty with a warning instead of stopping the import.
    """

    text = _text(value).lower()
    if not text:
        return "предложение", None
    if text not in _PRICE_TYPES:
        return "", (
            f"Тип цены «{_text(value)}» не распознан (ожидается «сделка» или "
            "«предложение»): поле оставлено пустым."
        )
    return _PRICE_TYPES[text], None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _default_collected_at() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_listing(
    row: Mapping[str, Any],
    source: Optional[str] = None,
    listing_type: str = "property",
    collected_at: Optional[str] = None,
    rent_period: Optional[str] = None,
) -> Dict[str, Any]:
    """Normalize one marketplace row into the shared listing schema.

    ``date`` is the publication date of the offer (the price date);
    ``collected_at`` is when the file was collected and is not a price date.
    ``price_type`` defaults to "предложение": marketplace listings are offers.

    ``rent``: the rent of the premises (``rent_rub``, or the price column) or
    the rate per m² (``rent_sqm``, with the area) per ``rent_period``
    (``month`` or ``year``, from the column «период» or the argument; no
    default) becomes the annual rent in ``price_rub``. ``income`` needs
    ``noi`` and/or ``gross_income``; ``business`` needs ``value`` (or the
    price column) and ``metric``; ``machinery`` reads ``name`` and
    ``operating_hours``.
    """

    if not isinstance(row, Mapping):
        raise ValueError("row must be an object")
    listing_kind = _text(listing_type).lower()
    if listing_kind not in LISTING_TYPES:
        raise ValueError(f"listing_type must be one of {', '.join(sorted(LISTING_TYPES))}")

    source_name = _text(source or _lookup(row, "source"))
    if source:
        source_name = _text(source)
    if not source_name:
        raise ValueError("source must be provided")

    values: Dict[str, Any] = {
        "listing_type": listing_kind,
        "source": source_name,
    }
    for field in _TEXT_FIELDS:
        values[field] = _text(_lookup(row, field))
    for field in _NUMBER_FIELDS:
        values[field] = parse_number(_lookup(row, field), field)
    for field in ("adjustment_pct", "weight"):
        number = parse_number(_lookup(row, field), field)
        if number is not None:
            values[field] = number
    for field in _TYPE_TEXT_FIELDS.get(listing_kind, ()):
        values[field] = _text(_lookup(row, field))
    for field in _TYPE_NUMBER_FIELDS.get(listing_kind, ()):
        values[field] = parse_number(_lookup(row, field), field)
    if listing_kind == "rent":
        _annual_rent(row, values, rent_period)
    if listing_kind == "business":
        if values["value"] is None:
            values["value"] = values["price_rub"]
        if values["value"] is None or values["value"] <= 0:
            raise ValueError("value (стоимость компании) must be a number greater than 0")
        if values["metric"] is None or values["metric"] <= 0:
            raise ValueError("metric (показатель) must be a number greater than 0")
    if listing_kind == "income" and values["noi"] is None and values["gross_income"] is None:
        raise ValueError("noi (ЧОД) or gross_income (валовой доход) is required for income listings")
    if listing_kind == "machinery" and values["operating_hours"] is not None and values["operating_hours"] < 0:
        raise ValueError("operating_hours must be non-negative")
    # A row normalized earlier keeps its warnings; an unknown price type
    # marked then stays empty instead of becoming an offer.
    previous = row.get("import_warnings")
    warnings = [str(item) for item in previous] if isinstance(previous, list) else []
    raw_price_type = _lookup(row, "price_type")
    if warnings and not _text(raw_price_type):
        values["price_type"], price_type_warning = "", None
    else:
        values["price_type"], price_type_warning = _price_type(raw_price_type)
    if price_type_warning:
        warnings.append(price_type_warning)
    if warnings:
        values["import_warnings"] = warnings
    adjustments = _parse_adjustments(_lookup(row, "adjustments"))
    if adjustments is not None:
        values["adjustments"] = adjustments

    if collected_at is not None:
        values["collected_at"] = _text(collected_at)
    elif not values["collected_at"]:
        values["collected_at"] = _default_collected_at()

    if listing_kind != "business" and (values["price_rub"] is None or values["price_rub"] <= 0):
        raise ValueError("price_rub must be a number greater than 0")
    if values["area_sqm"] is not None and values["area_sqm"] <= 0:
        raise ValueError("area_sqm must be greater than 0 when provided")
    if values["mileage_km"] is not None and values["mileage_km"] < 0:
        raise ValueError("mileage_km must be non-negative")

    if values["listing_id"]:
        values["listing_id_basis"] = "listing_id"
    else:
        # Without url, VIN or cadastral number the id is built from the
        # characteristics. The address alone is not an identifier (flats
        # share it), but it distinguishes listings together with the rest.
        if values["url"]:
            basis, stable_fields = "url", (values["source"], values["url"])
        elif values["vin"]:
            basis, stable_fields = "vin", (values["source"], "vin", values["vin"].upper())
        elif values["cadastral_number"]:
            basis, stable_fields = "cadastral", (values["source"], "cadastral", values["cadastral_number"])
        else:
            basis = "characteristics"
            stable_fields = tuple(
                values.get(field)
                for field in (
                    "source", "listing_type", "price_rub", "area_sqm", "address", "city",
                    "rooms", "floor", "total_floors", "year", "brand", "model", "mileage_km",
                    "name", "value", "metric", "operating_hours",
                )
            )
        seed = "|".join("" if item is None else str(item) for item in stable_fields)
        values["listing_id"] = hashlib.sha256(seed.encode("utf-8")).hexdigest()[:16]
        values["listing_id_basis"] = basis

    return values


def _annual_rent(row: Mapping[str, Any], values: Dict[str, Any], rent_period: Optional[str]) -> None:
    """Annual rent of the premises into ``price_rub``; the period is never assumed."""

    amount = parse_number(_lookup(row, "rent_rub"), "rent_rub")
    if amount is None:
        amount = values["price_rub"]
    rate = parse_number(_lookup(row, "rent_sqm"), "rent_sqm")
    raw_period = _text(_lookup(row, "rent_period")) or _text(rent_period)
    period = _RENT_PERIODS.get(raw_period.lower())
    if period is None:
        raise ValueError(
            "rent_period must be 'month' or 'year' (column «период» or the rent_period argument)"
        )
    if amount is None and rate is None:
        raise ValueError("rent_rub (арендная плата) or rent_sqm (ставка за м²) is required for rent listings")
    if amount is None:
        if values["area_sqm"] is None:
            raise ValueError("area_sqm is required with rent_sqm")
        amount = rate * values["area_sqm"]
    if amount <= 0:
        raise ValueError("rent must be greater than 0")
    values["rent_rub"] = amount
    values["rent_sqm"] = rate
    values["rent_period"] = period
    values["price_rub"] = amount * (12 if period == "month" else 1)


def deduplicate_listings(listings: Iterable[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    """Keep the first row for each source/listing_id pair."""

    result: List[Dict[str, Any]] = []
    seen = set()
    for listing in listings:
        key = (listing.get("source"), listing.get("listing_id"))
        if key in seen:
            continue
        seen.add(key)
        result.append(dict(listing))
    return result


def _json_rows(path: Path) -> Iterable[Tuple[str, Any]]:
    with path.open("r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if isinstance(payload, Mapping):
        for key in ("items", "listings", "data"):
            candidate = payload.get(key)
            if isinstance(candidate, list):
                payload = candidate
                break
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON root must be an object or an array")
    for number, row in enumerate(payload, start=1):
        yield f"item {number}", row


def _jsonl_rows(path: Path) -> Iterable[Tuple[str, Mapping[str, Any]]]:
    with path.open("r", encoding="utf-8-sig") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"invalid JSONL at line {line_number}") from exc
            if not isinstance(row, Mapping):
                raise ValueError(f"JSONL line {line_number} must contain an object")
            yield f"line {line_number}", row


def _csv_rows(path: Path) -> Iterable[Tuple[str, Mapping[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        # Sniff the delimiter from the header only: data cells may hold JSON
        # with commas. Quoting follows standard CSV ("" inside quoted cells).
        sample = handle.readline()
        handle.seek(0)
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.DictReader(handle, delimiter=delimiter)
        if not reader.fieldnames:
            raise ValueError("CSV must contain a header row")
        for row in reader:
            yield f"line {reader.line_num}", row


def _xlsx_rows(path: Path, sheet: Optional[str]) -> Iterable[Tuple[str, Mapping[str, Any]]]:
    """Rows of an Excel sheet; the header is the first non-empty row."""

    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency of the package
        raise ValueError("reading .xlsx needs the openpyxl package") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            worksheet = workbook.worksheets[0]
        elif sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            raise ValueError(f"sheet «{sheet}» not found; sheets: {', '.join(workbook.sheetnames)}")
        header: Optional[List[str]] = None
        for number, cells in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if all(cell in (None, "") for cell in cells):
                continue
            if header is None:
                header = ["" if cell is None else str(cell).strip() for cell in cells]
                continue
            row = {}
            for key, cell in zip(header, cells):
                if not key:
                    continue
                if isinstance(cell, datetime):
                    cell = cell.date().isoformat() if cell.time() == datetime.min.time() else cell.isoformat()
                row[key] = cell
            yield f"row {number}", row
        if header is None:
            raise ValueError("the sheet must contain a header row")
    finally:
        workbook.close()


def load_listings(
    path: str,
    source: str,
    listing_type: str = "property",
    collected_at: Optional[str] = None,
    sheet: Optional[str] = None,
    rent_period: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Load and normalize a local CSV, JSON, JSONL or Excel (.xlsx) export.

    ``sheet`` names the Excel sheet (the first one by default); the header is
    the first non-empty row. ``rent_period`` applies to ``rent`` listings
    without a period column.

    No network request is made. Duplicate source/listing_id rows are removed
    while preserving the first occurrence. An invalid row stops the import
    with its line (CSV, JSONL) or item (JSON) number in the error.
    """

    input_path = Path(path)
    if not input_path.is_file():
        raise FileNotFoundError(str(input_path))
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        rows = _csv_rows(input_path)
    elif suffix == ".json":
        rows = _json_rows(input_path)
    elif suffix == ".jsonl":
        rows = _jsonl_rows(input_path)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _xlsx_rows(input_path, sheet)
    else:
        raise ValueError("supported input formats are .csv, .json, .jsonl and .xlsx")

    normalized = []
    for label, row in rows:
        try:
            normalized.append(normalize_listing(row, source, listing_type, collected_at, rent_period))
        except ValueError as exc:
            raise ValueError(f"{label}: {exc}") from exc
    return deduplicate_listings(normalized)
